"""Save a register workbook, then reload and verify every geometry cell. Refuses silently-truncated output."""
import openpyxl
from shapely import wkt as W
GEOM_COLS = ['Basin outline (WKT) - surveyed',
             'Basin footprint - DERIVED (convex hull of real field points, NOT surveyed)',
             'Basin footprint - APPROXIMATE (user-supplied bounding box, NOT surveyed or derived)',
             'Field outline (WKT)', 'Field outline (WKT, from OGIM v3.0)']
def safe_save(wb, path):
    wb.save(path)
    chk = openpyxl.load_workbook(path, data_only=True)
    problems, n = [], 0
    for ws in chk.worksheets:
        hdr = {c: ws.cell(row=1,column=c).value for c in range(1, ws.max_column+1)}
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                v = cell.value
                if isinstance(v, str) and len(v) >= 32767:
                    problems.append(f"{ws.title}!{cell.coordinate} at 32,767-char cell limit (truncated)")
                if hdr.get(cell.column) in GEOM_COLS and v:
                    n += 1
                    try: W.loads(v)
                    except Exception as e: problems.append(f"{ws.title}!{cell.coordinate} geometry does not parse: {str(e)[:60]}")
    if problems:
        raise RuntimeError("SAVE VERIFICATION FAILED:\n  " + "\n  ".join(problems[:30]))
    print(f"safe_save OK: {path} | {n} geometry cells all parse | no cell at the length limit")
